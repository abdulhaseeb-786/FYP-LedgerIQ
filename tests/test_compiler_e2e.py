"""End-to-end integration and dynamic formula verification tests."""

import openpyxl
from pathlib import Path

SAMPLE_DIR = Path(__file__).parent.parent / "data" / "sample_tbs"
from fyp_module3.cli import run_pipeline


def test_e2e_compiler_tb1(tmp_path):
    tb1_path = SAMPLE_DIR / "TB 1.xlsx"

    out_file = str(tmp_path / "TB1_Compiled_FS.xlsx")
    generated_path = run_pipeline(str(tb1_path), out_file, entity_name="Grant Project 1")
    assert Path(generated_path).exists()

    # Inspect with openpyxl (formulas retained)
    wb = openpyxl.load_workbook(generated_path, data_only=False)
    expected_sheets = ["TB", "Balance Sheet", "Profit or Loss", "Changes in Equity", "Cash Flows", "Notes", "Mapping Audit"]
    for s in expected_sheets:
        assert s in wb.sheetnames, f"Missing sheet: {s}"

    # Verify Notes Sheet formulas
    ws_notes = wb["Notes"]
    found_sum = False
    for row in ws_notes.iter_rows(values_only=True):
        for cell in row:
            if cell and str(cell).startswith("=SUM("):
                found_sum = True
                break
    assert found_sum is True, "Notes sheet must contain dynamic =SUM() formulas"

    # Verify P&L Formulas
    ws_pl = wb["Profit or Loss"]
    found_pl_formula = False
    for row in ws_pl.iter_rows(values_only=True):
        for cell in row:
            if cell and (str(cell).startswith("=") or "'Notes'!" in str(cell)):
                found_pl_formula = True
                break
    assert found_pl_formula is True, "Profit or Loss must contain dynamic formulas"

    # Verify Balance Sheet Check formula
    ws_bs = wb["Balance Sheet"]
    found_bs_check = False
    for row in ws_bs.iter_rows(values_only=True):
        for cell in row:
            if cell and str(cell).startswith("=C") and "-" in str(cell):
                found_bs_check = True
                break
    assert found_bs_check is True, "Balance Sheet must contain dynamic formula check"


def test_e2e_compiler_tb2(tmp_path):
    tb2_path = SAMPLE_DIR / "TB 2.xlsx"

    out_file = str(tmp_path / "TB2_Compiled_FS.xlsx")
    generated_path = run_pipeline(str(tb2_path), out_file, entity_name="Services Enterprise")
    assert Path(generated_path).exists()

    wb = openpyxl.load_workbook(generated_path, data_only=False)
    assert len(wb.sheetnames) == 7

    ws_bs = wb["Balance Sheet"]
    # Verify no self-referential / circular formulas exist
    for r in range(1, ws_bs.max_row + 1):
        cell_formula = ws_bs.cell(row=r, column=3).value
        if cell_formula and str(cell_formula).startswith("="):
            # A formula on row R should never sum C{R}:C{R}
            assert f"C{r}:C{r}" not in str(cell_formula), f"Circular reference detected at row {r}: {cell_formula}"
            assert f"C{r}+" not in str(cell_formula) and f"+C{r}" not in str(cell_formula), f"Circular reference at row {r}"

