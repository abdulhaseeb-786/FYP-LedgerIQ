"""Post-export audit of the formulas actually written into the workbook.

The compilers build two representations of every figure: a Python amount and
an Excel formula. The workbook ships the formula. Verifying the amount and
then exporting the formula is how a run can report "100% balanced" while the
delivered file is out by the value of a whole account.

This module closes that gap by re-reading the saved file and evaluating its
formulas, so the integrity claim is made about the artefact the user receives.
"""

import re
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from fyp_module3.core.models import WorkbookAuditResult

TOLERANCE = 0.01

_SUMIF = re.compile(
    r"SUMIF\(\s*"
    r"(?:'(?P<ksheet>[^']+)'|(?P<ksheet2>\w+))!"
    r"\$?(?P<kcol>[A-Z]{1,2})\$?(?P<krow1>\d*):\$?[A-Z]{1,2}\$?(?P<krow2>\d*)\s*,"
    r"\s*\"(?P<criterion>[^\"]*)\"\s*,"
    r"\s*(?:'(?P<vsheet>[^']+)'|(?P<vsheet2>\w+))!"
    r"\$?(?P<vcol>[A-Z]{1,2})\$?\d*:\$?[A-Z]{1,2}\$?\d*\s*\)",
    re.IGNORECASE,
)
_SUM = re.compile(
    r"SUM\(\s*(?:(?:'([^']+)'|(\w+))!)?\$?([A-Z]{1,2})\$?(\d+)\s*:\s*"
    r"(?:(?:'([^']+)'|(\w+))!)?\$?([A-Z]{1,2})\$?(\d+)\s*\)",
    re.IGNORECASE,
)
_REF = re.compile(r"(?:('[^']+'|\b[A-Za-z_][\w.]*)!)?\$?([A-Z]{1,2})\$?(\d+)\b")


class FormulaEvaluator:
    """Evaluates the subset of Excel used by this compiler."""

    def __init__(self, workbook):
        self.wb = workbook
        self._cache: Dict[Tuple[str, str], float] = {}
        self._in_progress: set = set()

    # -----------------------------------------------------------------
    def cell(self, sheet: str, column: str, row: int) -> float:
        key = (sheet, f"{column}{row}")
        if key in self._cache:
            return self._cache[key]
        if key in self._in_progress:
            raise ValueError(f"Circular reference at {sheet}!{column}{row}")
        if sheet not in self.wb.sheetnames:
            return 0.0

        self._in_progress.add(key)
        try:
            value = self.evaluate(sheet, self.wb[sheet][f"{column}{row}"].value)
        finally:
            self._in_progress.discard(key)
        self._cache[key] = value
        return value

    def column_values(self, sheet: str, column: str) -> List[Any]:
        if sheet not in self.wb.sheetnames:
            return []
        ws = self.wb[sheet]
        return [ws[f"{column}{r}"].value for r in range(1, ws.max_row + 1)]

    # -----------------------------------------------------------------
    def evaluate(self, sheet: str, raw: Any) -> float:
        if raw is None or raw == "":
            return 0.0
        if isinstance(raw, bool):
            return 0.0
        if isinstance(raw, (int, float)):
            return float(raw)

        text = str(raw)
        if not text.startswith("="):
            try:
                return float(text)
            except ValueError:
                return 0.0

        expr = text[1:]
        expr = _SUMIF.sub(lambda m: repr(self._eval_sumif(sheet, m)), expr)
        expr = _SUM.sub(lambda m: repr(self._eval_sum(sheet, m)), expr)
        expr = _REF.sub(lambda m: repr(self._eval_ref(sheet, m)), expr)

        if not re.fullmatch(r"[-+*/().\de\s]*", expr):
            raise ValueError(f"Unsupported formula on sheet '{sheet}': {text}")
        if not expr.strip():
            return 0.0
        return float(eval(expr))  # noqa: S307 - expression is validated above

    # -----------------------------------------------------------------
    def _eval_sumif(self, sheet: str, m: re.Match) -> float:
        crit_sheet = m.group("ksheet") or m.group("ksheet2") or sheet
        sum_sheet = m.group("vsheet") or m.group("vsheet2") or sheet
        crit_col, sum_col = m.group("kcol"), m.group("vcol")
        criterion = m.group("criterion")

        if crit_sheet not in self.wb.sheetnames:
            return 0.0
        first = int(m.group("krow1")) if m.group("krow1") else 1
        last = int(m.group("krow2")) if m.group("krow2") else self.wb[crit_sheet].max_row

        total = 0.0
        for row in range(min(first, last), max(first, last) + 1):
            key = self.wb[crit_sheet][f"{crit_col}{row}"].value
            if key is None or str(key).strip() != criterion:
                continue
            total += self.cell(sum_sheet, sum_col, row)
        return round(total, 6)

    def _eval_sum(self, sheet: str, m: re.Match) -> float:
        target = m.group(1) or m.group(2) or sheet
        col, r1 = m.group(3), int(m.group(4))
        r2 = int(m.group(8))
        return round(sum(self.cell(target, col, r) for r in range(min(r1, r2), max(r1, r2) + 1)), 6)

    def _eval_ref(self, sheet: str, m: re.Match) -> float:
        target = (m.group(1) or sheet).strip("'")
        return self.cell(target, m.group(2), int(m.group(3)))


# ---------------------------------------------------------------------
def _find_row(ws, text: str, column: int = 1) -> Optional[int]:
    needle = text.lower()
    for r in range(1, ws.max_row + 1):
        value = ws.cell(row=r, column=column).value
        if value and needle in str(value).lower():
            return r
    return None


def audit_workbook(path: str) -> WorkbookAuditResult:
    """Evaluate the saved workbook and confirm it states what it claims."""
    wb = openpyxl.load_workbook(path, data_only=False)
    ev = FormulaEvaluator(wb)
    result = WorkbookAuditResult()

    def record(ok: bool, label: str, detail: str) -> None:
        if ok:
            result.checks.append(f"{label}: {detail}")
        else:
            result.passed = False
            result.failures.append(f"{label}: {detail}")

    # 1. Trial balance parity, as the workbook computes it.
    try:
        ws = wb["TB"]
        row = _find_row(ws, "Checksum Check", column=2)
        variance = ev.cell("TB", "E", row) if row else 0.0
        result.trial_balance_variance = round(variance, 2)
        record(abs(variance) < TOLERANCE, "Trial balance parity",
               f"debits less credits = {variance:,.2f}")
    except Exception as exc:                                   # pragma: no cover
        record(False, "Trial balance parity", f"could not evaluate ({exc})")

    # 2. The balance sheet identity, as the workbook computes it.
    try:
        ws = wb["Balance Sheet"]
        assets_row = _find_row(ws, "TOTAL ASSETS")
        eql_row = _find_row(ws, "TOTAL EQUITY AND LIABILITIES")
        assets = ev.cell("Balance Sheet", "C", assets_row)
        eql = ev.cell("Balance Sheet", "C", eql_row)
        variance = round(assets - eql, 2)
        result.balance_sheet_variance = variance
        record(abs(variance) < TOLERANCE, "Balance sheet identity (A = L + E)",
               f"assets {assets:,.2f} less equity and liabilities {eql:,.2f} = {variance:,.2f}")
    except Exception as exc:                                   # pragma: no cover
        record(False, "Balance sheet identity (A = L + E)", f"could not evaluate ({exc})")

    # 3. The cash flow statement must close on the balance sheet's cash.
    try:
        ws = wb["Cash Flows"]
        row = _find_row(ws, "Variance")
        variance = round(ev.cell("Cash Flows", "C", row), 2) if row else 0.0
        result.cash_flow_variance = variance
        record(abs(variance) < TOLERANCE, "Cash flow reconciles to balance sheet",
               f"closing cash variance = {variance:,.2f}")
    except Exception as exc:                                   # pragma: no cover
        record(False, "Cash flow reconciles to balance sheet", f"could not evaluate ({exc})")

    # 4. The equity statement must close on the balance sheet's equity.
    try:
        eq_ws = wb["Changes in Equity"]
        bs_ws = wb["Balance Sheet"]
        eq_row = _find_row(eq_ws, "Balance as at")
        bs_row = _find_row(bs_ws, "Total Equity")
        closing = ev.cell("Changes in Equity", "D", eq_row)
        bs_equity = ev.cell("Balance Sheet", "C", bs_row)
        variance = round(closing - bs_equity, 2)
        result.equity_variance = variance
        record(abs(variance) < TOLERANCE, "Equity statement ties to balance sheet",
               f"closing equity {closing:,.2f} against {bs_equity:,.2f} = {variance:,.2f}")
    except Exception as exc:                                   # pragma: no cover
        record(False, "Equity statement ties to balance sheet", f"could not evaluate ({exc})")

    # 5. Every note subtotal must equal the face line that references it.
    try:
        linked = 0
        for sheet_name in ("Balance Sheet", "Profit or Loss"):
            ws = wb[sheet_name]
            for r in range(1, ws.max_row + 1):
                formula = ws.cell(row=r, column=3).value
                if not isinstance(formula, str) or "'Notes'!" not in formula:
                    continue
                linked += 1
                face = ev.cell(sheet_name, "C", r)
                note_cell = re.search(r"'Notes'!\$?([A-Z]{1,2})\$?(\d+)", formula)
                note_value = ev.cell("Notes", note_cell.group(1), int(note_cell.group(2)))
                label = ws.cell(row=r, column=1).value
                if abs(face - note_value) >= TOLERANCE:
                    record(False, "Note cross-reference",
                           f"'{label}' shows {face:,.2f} but its note totals {note_value:,.2f}")
        record(True, "Note cross-references", f"{linked} face line(s) linked to note subtotals")
    except Exception as exc:                                   # pragma: no cover
        record(False, "Note cross-references", f"could not evaluate ({exc})")

    return result
