"""Corporate accounting styles, borders, and number formats for OpenXML workbooks."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

FONT_FAMILY = "Segoe UI"

# Typography
FONT_COMPANY = Font(name=FONT_FAMILY, size=13, bold=True, color="1E293B")
FONT_TITLE = Font(name=FONT_FAMILY, size=11, bold=True, color="0F172A")
FONT_SUBTITLE = Font(name=FONT_FAMILY, size=9.5, italic=True, color="64748B")
FONT_TH = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name=FONT_FAMILY, size=10, color="1E293B")
FONT_BOLD = Font(name=FONT_FAMILY, size=10, bold=True, color="0F172A")
FONT_SECTION = Font(name=FONT_FAMILY, size=10.5, bold=True, color="1E3A8A")
FONT_CHECK_PASS = Font(name=FONT_FAMILY, size=9.5, bold=True, color="16A34A")
FONT_ALERT_HIGH = Font(name=FONT_FAMILY, size=9.5, bold=True, color="B91C1C")
FONT_ALERT_MED = Font(name=FONT_FAMILY, size=9.5, bold=True, color="D97706")

# Fills
FILL_TH = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
FILL_SUBTOTAL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
FILL_SECTION = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
FILL_ALERT_HIGH = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
FILL_ALERT_MED = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
FILL_ALERT_PASS = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

# Borders
THIN_LINE = Side(style="thin", color="94A3B8")
DOUBLE_LINE = Side(style="double", color="0F172A")

BORDER_TH = Border(top=THIN_LINE, bottom=THIN_LINE, left=THIN_LINE, right=THIN_LINE)
BORDER_SUBTOTAL = Border(top=THIN_LINE, bottom=THIN_LINE)
BORDER_GRAND_TOTAL = Border(top=THIN_LINE, bottom=DOUBLE_LINE)
BORDER_CHECK = Border(top=THIN_LINE, bottom=THIN_LINE, left=THIN_LINE, right=THIN_LINE)

# Alignments
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

# Standard Accounting Number Format Mask
ACCOUNTING_NUMBER_FORMAT = '_(* #,##0.00_);_(* (#,##0.00);_(* "-"??_);_(@_)'
INTEGER_NUMBER_FORMAT = '#,##0'
PERCENTAGE_FORMAT = '0.0%'


def autofit_columns(ws, max_len_cap: int = 65):
    """Adjust worksheet column widths to fit content with padding."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = cell.value
            if val is not None:
                # If formula string, estimate standard width
                s = str(val)
                if s.startswith("="):
                    s = " 1,234,567.89 "
                max_len = max(max_len, len(s))
        ws.column_dimensions[col_letter].width = min(max_len_cap, max(12, max_len + 3))

