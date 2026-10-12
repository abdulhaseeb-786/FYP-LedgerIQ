"""Multi-Tab Dynamic OpenXML Workbook Compiler (Module 9 Exporter)."""

from typing import List
import openpyxl

from fyp_module3.core.models import StandardizedAccountRecord, AuditEntry, CompiledNote
from fyp_module3.statements.notes_compiler import NotesCompiler
from fyp_module3.statements.profit_loss import ProfitLossStatement
from fyp_module3.statements.balance_sheet import BalanceSheetStatement
from fyp_module3.statements.equity import EquityStatement
from fyp_module3.statements.cash_flow import CashFlowStatement
from fyp_module3.statements.formulas import tb_row_amount
from fyp_module3.exporter.styles import (
    FONT_COMPANY, FONT_TITLE, FONT_SUBTITLE, FONT_TH, FONT_REGULAR, FONT_BOLD,
    FONT_SECTION, FILL_TH, FILL_SECTION,
    BORDER_TH, BORDER_SUBTOTAL, BORDER_GRAND_TOTAL, BORDER_CHECK,
    ALIGN_LEFT, ALIGN_CENTER, ALIGN_RIGHT, ACCOUNTING_NUMBER_FORMAT,
    autofit_columns
)


class ExcelWorkbookCompiler:
    """Assembles and writes the full 7-tab OpenXML financial statement package."""

    def __init__(self, entity_name: str = "Client Company Ltd.", period: str = "June 30, 2025"):
        self.entity_name = entity_name
        self.period = period

    def export_workbook(
        self,
        accounts: List[StandardizedAccountRecord],
        audit_trail: List[AuditEntry],
        output_file: str
    ) -> str:
        """Compile and save the multi-tab linked workbook."""
        wb = openpyxl.Workbook()
        # Remove default empty sheet
        wb.remove(wb.active)

        # Accounts are written to the TB sheet from this row down, and the
        # notes point their formulas back at those rows.
        tb_start_row = 6

        # 1. Compile Notes
        notes_compiler = NotesCompiler()
        compiled_notes = notes_compiler.compile_notes(
            accounts, start_row=6, tb_start_row=tb_start_row)

        # 2. Compile Profit or Loss
        pl_compiler = ProfitLossStatement(notes_compiler)
        pl_rows = pl_compiler.compile(accounts, start_row=6, tb_first_row=tb_start_row)

        # 3. Compile Balance Sheet
        bs_compiler = BalanceSheetStatement(notes_compiler)
        bs_rows = bs_compiler.compile(
            accounts,
            net_profit_ref=pl_compiler.net_profit_cell_ref,
            net_profit=pl_compiler.net_profit,
            start_row=6,
            tb_first_row=tb_start_row,
        )

        # 4. Compile Equity
        eq_compiler = EquityStatement()
        eq_rows = eq_compiler.compile(
            accounts,
            net_profit_ref=pl_compiler.net_profit_cell_ref,
            net_profit=pl_compiler.net_profit,
            period_label=self.period,
            start_row=6,
            tb_first_row=tb_start_row,
        )

        # 5. Compile Cash Flow
        cf_compiler = CashFlowStatement()
        cf_rows = cf_compiler.compile(
            accounts,
            pbt_ref=pl_compiler.pbt_cell_ref,
            cash_note_ref=notes_compiler.get_note_cell_ref(8),
            profit_before_tax=pl_compiler.profit_before_tax,
            start_row=6,
            tb_first_row=tb_start_row,
        )

        # Published so the caller can audit what was actually written.
        self.notes_compiler = notes_compiler
        self.pl_compiler = pl_compiler
        self.bs_compiler = bs_compiler
        self.eq_compiler = eq_compiler
        self.cf_compiler = cf_compiler
        self.reconciliation_errors = (list(bs_compiler.reconciliation_errors)
                                      + list(pl_compiler.reconciliation_errors))

        # Build each worksheet
        self._write_tb_sheet(wb, accounts)
        self._write_balance_sheet(wb, bs_rows)
        self._write_profit_loss(wb, pl_rows)
        self._write_equity_sheet(wb, eq_rows)
        self._write_cash_flow_sheet(wb, cf_rows)
        self._write_notes_sheet(wb, compiled_notes)
        self._write_audit_sheet(wb, accounts, audit_trail)
        # Force Excel to recalculate all formula cells on file open
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalculation = True

        wb.save(output_file)
        return output_file

    def _write_sheet_header(self, ws, title: str):
        """Write uniform corporate statement header."""
        ws.cell(row=1, column=1, value=self.entity_name).font = FONT_COMPANY
        ws.cell(row=2, column=1, value=title).font = FONT_TITLE
        ws.cell(row=3, column=1, value=f"FOR THE PERIOD ENDED {self.period.upper()}").font = FONT_SUBTITLE

    def _write_tb_sheet(self, wb, accounts: List[StandardizedAccountRecord]):
        ws = wb.create_sheet(title="TB")
        self._write_sheet_header(ws, "WORKING TRIAL BALANCE (BALANCED)")

        headers = ["Account Code", "Account Title", "Classification Line", "Note", "Debit (PKR)", "Credit (PKR)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_CENTER if col_idx in (1, 4) else (ALIGN_RIGHT if col_idx in (5, 6) else ALIGN_LEFT)
            cell.border = BORDER_TH

        row_idx = 6
        for acc in accounts:
            ws.cell(row=row_idx, column=1, value=acc.account_code or "").alignment = ALIGN_CENTER
            ws.cell(row=row_idx, column=2, value=acc.account_name).alignment = ALIGN_LEFT
            ws.cell(row=row_idx, column=3, value=acc.line_item_id or "").alignment = ALIGN_LEFT
            ws.cell(row=row_idx, column=4, value=acc.note_number if acc.note_number else "").alignment = ALIGN_CENTER
            
            c_dr = ws.cell(row=row_idx, column=5, value=acc.debit)
            c_dr.number_format = ACCOUNTING_NUMBER_FORMAT
            c_dr.alignment = ALIGN_RIGHT
            
            c_cr = ws.cell(row=row_idx, column=6, value=acc.credit)
            c_cr.number_format = ACCOUNTING_NUMBER_FORMAT
            c_cr.alignment = ALIGN_RIGHT
            row_idx += 1

        # Total Row with dynamic formulas
        ws.cell(row=row_idx, column=2, value="Total Working Trial Balance").font = FONT_BOLD
        total_dr_cell = ws.cell(row=row_idx, column=5, value=f"=SUM(E6:E{row_idx-1})")
        total_dr_cell.font = FONT_BOLD
        total_dr_cell.number_format = ACCOUNTING_NUMBER_FORMAT
        total_dr_cell.border = BORDER_GRAND_TOTAL

        total_cr_cell = ws.cell(row=row_idx, column=6, value=f"=SUM(F6:F{row_idx-1})")
        total_cr_cell.font = FONT_BOLD
        total_cr_cell.number_format = ACCOUNTING_NUMBER_FORMAT
        total_cr_cell.border = BORDER_GRAND_TOTAL

        # Parity Check row
        row_idx += 1
        ws.cell(row=row_idx, column=2, value="Checksum Check (Debit - Credit)").font = FONT_SUBTITLE
        diff_cell = ws.cell(row=row_idx, column=5, value=f"=E{row_idx-1}-F{row_idx-1}")
        diff_cell.font = FONT_BOLD
        diff_cell.number_format = ACCOUNTING_NUMBER_FORMAT

        autofit_columns(ws)

    def _write_balance_sheet(self, wb, rows: List[dict]):
        ws = wb.create_sheet(title="Balance Sheet")
        self._write_sheet_header(ws, "STATEMENT OF FINANCIAL POSITION")

        headers = ["Statement Item", "Note", "2025 (PKR)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_CENTER if col_idx == 2 else (ALIGN_RIGHT if col_idx == 3 else ALIGN_LEFT)
            cell.border = BORDER_TH

        for item in rows:
            r = item["row"]
            label_cell = ws.cell(row=r, column=1, value=item["label"])
            note_cell = ws.cell(row=r, column=2, value=item["note"] or "")
            amt_cell = ws.cell(row=r, column=3)

            note_cell.alignment = ALIGN_CENTER

            style = item.get("style", "normal")
            if style == "section_header":
                label_cell.font = FONT_SECTION
                label_cell.fill = FILL_SECTION
            elif style == "sub_header":
                label_cell.font = FONT_BOLD
            elif style == "subtotal":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_SUBTOTAL
            elif style == "grand_total":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_GRAND_TOTAL
            elif style == "check":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_CHECK
            else:
                label_cell.font = FONT_REGULAR

            # Assign formula or value
            if item["formula"]:
                amt_cell.value = item["formula"]
            elif item["amount"] is not None:
                amt_cell.value = item["amount"]

            amt_cell.number_format = ACCOUNTING_NUMBER_FORMAT
            amt_cell.alignment = ALIGN_RIGHT

        autofit_columns(ws)

    def _write_profit_loss(self, wb, rows: List[dict]):
        ws = wb.create_sheet(title="Profit or Loss")
        self._write_sheet_header(ws, "STATEMENT OF PROFIT OR LOSS")

        headers = ["Particulars", "Note", "2025 (PKR)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_CENTER if col_idx == 2 else (ALIGN_RIGHT if col_idx == 3 else ALIGN_LEFT)
            cell.border = BORDER_TH

        for item in rows:
            r = item["row"]
            label_cell = ws.cell(row=r, column=1, value=item["label"])
            note_cell = ws.cell(row=r, column=2, value=item["note"] or "")
            amt_cell = ws.cell(row=r, column=3)

            note_cell.alignment = ALIGN_CENTER

            style = item.get("style", "normal")
            if style == "subtotal":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_SUBTOTAL
            elif style == "grand_total":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_GRAND_TOTAL
            else:
                label_cell.font = FONT_REGULAR

            if item["formula"]:
                amt_cell.value = item["formula"]
            else:
                amt_cell.value = item["amount"]

            amt_cell.number_format = ACCOUNTING_NUMBER_FORMAT
            amt_cell.alignment = ALIGN_RIGHT

        autofit_columns(ws)

    def _write_equity_sheet(self, wb, rows: List[dict]):
        ws = wb.create_sheet(title="Changes in Equity")
        self._write_sheet_header(ws, "STATEMENT OF CHANGES IN EQUITY")

        headers = ["Particulars", "Share Capital / Capital", "Unappropriated Profit", "Total (PKR)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_RIGHT if col_idx > 1 else ALIGN_LEFT
            cell.border = BORDER_TH

        for item in rows:
            r = item["row"]
            label_cell = ws.cell(row=r, column=1, value=item["label"])
            cap_cell = ws.cell(row=r, column=2)
            re_cell = ws.cell(row=r, column=3)
            tot_cell = ws.cell(row=r, column=4, value=item["total_formula"])

            if "capital_formula" in item:
                cap_cell.value = item["capital_formula"]
            else:
                cap_cell.value = item["capital"]

            if "retained_formula" in item:
                re_cell.value = item["retained_formula"]
            else:
                re_cell.value = item["retained"]

            for c in (cap_cell, re_cell, tot_cell):
                c.number_format = ACCOUNTING_NUMBER_FORMAT
                c.alignment = ALIGN_RIGHT

            if item.get("style") == "grand_total":
                label_cell.font = FONT_BOLD
                for c in (cap_cell, re_cell, tot_cell):
                    c.font = FONT_BOLD
                    c.border = BORDER_GRAND_TOTAL
            else:
                label_cell.font = FONT_REGULAR

        autofit_columns(ws)

    def _write_cash_flow_sheet(self, wb, rows: List[dict]):
        ws = wb.create_sheet(title="Cash Flows")
        self._write_sheet_header(ws, "STATEMENT OF CASH FLOWS (INDIRECT METHOD)")

        headers = ["Particulars", "Note", "2025 (PKR)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_RIGHT if col_idx == 3 else ALIGN_LEFT
            cell.border = BORDER_TH

        for item in rows:
            r = item["row"]
            if not item["label"] and item.get("formula") is None and item.get("amount") is None:
                continue
            label_cell = ws.cell(row=r, column=1, value=item["label"])
            amt_cell = ws.cell(row=r, column=3)

            style = item.get("style", "normal")
            if style == "section_header":
                label_cell.font = FONT_SECTION
                label_cell.fill = FILL_SECTION
            elif style == "sub_header":
                label_cell.font = FONT_BOLD
            elif style == "subtotal":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_SUBTOTAL
            elif style == "grand_total":
                label_cell.font = FONT_BOLD
                amt_cell.font = FONT_BOLD
                amt_cell.border = BORDER_GRAND_TOTAL
            else:
                label_cell.font = FONT_REGULAR

            if item["formula"]:
                amt_cell.value = item["formula"]
            elif item["amount"] is not None:
                amt_cell.value = item["amount"]

            amt_cell.number_format = ACCOUNTING_NUMBER_FORMAT
            amt_cell.alignment = ALIGN_RIGHT

        autofit_columns(ws)

    def _write_notes_sheet(self, wb, notes: List[CompiledNote]):
        ws = wb.create_sheet(title="Notes")
        self._write_sheet_header(ws, "NOTES TO THE FINANCIAL STATEMENTS")

        for note in notes:
            # Note Title Header
            ws.cell(row=note.start_row, column=1, value=f"{note.note_number}.  {note.title}").font = FONT_SECTION
            
            # Table Column Header
            th_row = note.start_row + 1
            ws.cell(row=th_row, column=1, value="Account Description").font = FONT_BOLD
            ws.cell(row=th_row, column=3, value="2025 (PKR)").font = FONT_BOLD
            ws.cell(row=th_row, column=3).alignment = ALIGN_RIGHT

            # Items
            first_item_row = th_row + 1
            for idx, itm in enumerate(note.items):
                r = first_item_row + idx
                ws.cell(row=r, column=1, value=itm.account_name).font = FONT_REGULAR
                # Pull the figure live from the TB sheet so that editing a raw
                # balance there cascades into the notes and both statements.
                if itm.tb_sheet_row and itm.line_item_id:
                    value = tb_row_amount(itm.line_item_id, itm.tb_sheet_row)
                else:
                    value = itm.amount
                amt_cell = ws.cell(row=r, column=3, value=value)
                amt_cell.font = FONT_REGULAR
                amt_cell.number_format = ACCOUNTING_NUMBER_FORMAT
                amt_cell.alignment = ALIGN_RIGHT

            # Subtotal Row
            sub_r = note.subtotal_row
            ws.cell(row=sub_r, column=1, value=f"Total {note.title.title()}").font = FONT_BOLD
            sub_cell = ws.cell(row=sub_r, column=3, value=f"=SUM(C{first_item_row}:C{sub_r-1})")
            sub_cell.font = FONT_BOLD
            sub_cell.number_format = ACCOUNTING_NUMBER_FORMAT
            sub_cell.border = BORDER_SUBTOTAL
            sub_cell.alignment = ALIGN_RIGHT

        autofit_columns(ws)

    def _write_audit_sheet(self, wb, accounts: List[StandardizedAccountRecord], audit_trail: List[AuditEntry]):
        ws = wb.create_sheet(title="Mapping Audit")
        self._write_sheet_header(ws, "CLASSIFICATION MAPPINGS & BALANCING AUDIT TRAIL")

        headers = ["Account Code", "Account Name", "Mapped Line Item", "Note", "Confidence", "Status", "Adjustment Note"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=5, column=col_idx, value=h)
            cell.font = FONT_TH
            cell.fill = FILL_TH
            cell.alignment = ALIGN_CENTER if col_idx in (1, 4, 5, 6) else ALIGN_LEFT
            cell.border = BORDER_TH

        r = 6
        for acc in accounts:
            ws.cell(row=r, column=1, value=acc.account_code or "").alignment = ALIGN_CENTER
            ws.cell(row=r, column=2, value=acc.account_name).alignment = ALIGN_LEFT
            ws.cell(row=r, column=3, value=acc.line_item_id or "").alignment = ALIGN_LEFT
            ws.cell(row=r, column=4, value=acc.note_number if acc.note_number else "").alignment = ALIGN_CENTER
            ws.cell(row=r, column=5, value=f"{acc.confidence*100:.0f}%").alignment = ALIGN_CENTER
            
            status_cell = ws.cell(row=r, column=6, value="Adjusted" if acc.is_adjusted else "Original")
            status_cell.alignment = ALIGN_CENTER
            if acc.is_adjusted:
                status_cell.font = FONT_BOLD
            
            ws.cell(row=r, column=7, value=acc.adjustment_reason or "Mapped autonomously").alignment = ALIGN_LEFT
            r += 1

        # Balancing Audit Entries Section
        if audit_trail:
            r += 2
            ws.cell(row=r, column=1, value="HEURISTIC BALANCING ACTIONS APPLIED").font = FONT_SECTION
            r += 1
            audit_headers = ["Action", "Heuristic", "Account", "Before (Dr / Cr)", "After (Dr / Cr)", "Explanation"]
            for col_idx, ah in enumerate(audit_headers, 1):
                cell = ws.cell(row=r, column=col_idx, value=ah)
                cell.font = FONT_BOLD
                cell.fill = FILL_SECTION
                cell.border = BORDER_TH
            r += 1
            for entry in audit_trail:
                ws.cell(row=r, column=1, value=entry.action).font = FONT_BOLD
                ws.cell(row=r, column=2, value=entry.heuristic)
                ws.cell(row=r, column=3, value=entry.account_name)
                ws.cell(row=r, column=4, value=f"{entry.before_dr:,.2f} / {entry.before_cr:,.2f}")
                ws.cell(row=r, column=5, value=f"{entry.after_dr:,.2f} / {entry.after_cr:,.2f}")
                ws.cell(row=r, column=6, value=entry.explanation)
                r += 1

        autofit_columns(ws)
