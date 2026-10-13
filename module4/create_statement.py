"""
Builds the bank reconciliation statement from the reviewed files and adds it as a
"Reconciliation" sheet in the same workbook.

USAGE
    python reconciliation_statement.py
    (processes every output/review_*.xlsx, after a human has checked the Review sheet)

INPUT (each review_<name>.xlsx, made by review.py)
    Review sheet   : only rows still marked "No" in Match? are used
    Balances sheet : B1 = bank closing balance, B2 = ledger closing balance

FLOW
    1. Read the "No" rows and the two closing balances
    2. Split the rows into three groups:
         - amount mismatch : both entries exist, the amounts differ
         - unmatched bank  : bank entry with no ledger entry
         - unmatched ledger: ledger entry with no bank entry
       Rows marked "No" for any other reason (no match, low score) are broken into
       their bank entry and ledger entry, and each is treated as unmatched.
    3. Turn each group into signed lines (ADD / DEDUCT) on one side of the statement
    4. Write the statement, the adjusted balances, the check, and the formatting

RULES (which side, and add or deduct)
    Ledger entry not in the bank statement -> goes on the BANK side
        ledger debit  (money received) -> ADD to bank balance
        ledger credit (money paid out) -> DEDUCT from bank balance
        e.g. deposits in transit, outstanding cheques

    Bank entry not in the ledger -> goes on the CASH BOOK side
        bank credit (money received)   -> ADD to ledger balance
        bank debit  (money taken out)  -> DEDUCT from ledger balance
        e.g. interest received, bank fees, bounced cheques

    Amount mismatch -> only the difference (bank amount - ledger amount) goes on the
        CASH BOOK side, with the same signs as above. The error is assumed to be on the
        accountant's side, not the bank's. Labelled "ERROR: <description>".

    Bank and ledger use opposite labels for the same money, so the Type column in the
    review sheet is read as the bank's label if the row has a bank entry, else the ledger's.

OUTPUT (Reconciliation sheet)
    Rows 1-3 : company name, "BANK RECONCILIATION", "AS AT <today>"
    Left     : BANK STATEMENT (columns A:B)    Right: CASH BOOK (columns D:E)
    Each side: unadjusted closing balance, the signed lines, then the adjusted closing
               balance (a live Excel formula: unadjusted + sum of lines)
    Check    : UNRECONCILED AMOUNT = adjusted bank - adjusted cash book. The cell turns
               green and says "Reconciled" when it rounds to 0.00, red and "Not reconciled"
               otherwise.
    An existing Reconciliation sheet is replaced on every run. Review and Balances are
    left untouched.

NOTES
    - The company name is hardcoded as "Company".
    - Close the file in Excel before running, or the save fails.
"""

from openpyxl import load_workbook
from datetime import date
from review import headers
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pathlib import Path

def row_sign(row):
    """
    For entries with no matches
    Returns +1 (add) or -1 (subtract) for one review row.

    Cash ledger entry not in the bank statement -> adjusts the BANK balance:
        credit in cash ledger = money paid out     -> subtracted from bank balance
        debit in cash ledger  = money received     -> added to bank balance

    Bank statement entry not in the cash ledger -> adjusts the CASH LEDGER balance:
        credit in bank statement = money received  -> added to cash ledger balance
        debit in bank statement  = money taken out -> subtracted from cash ledger balance

    The two books use opposite labels for the same money.
    The Type column holds the bank's label if the row has a bank entry, else the ledger's label.
    """
    col = {name: headers.index(name) for name in headers}   # column name -> position in the row
    sign = 1 if row[col["Type"]] == "debit" else -1         # credit in ledger subtracted from bankbalance
    if row[col["Bank Amount"]] != "-":                      # it's a bank entry instead, so Type is the bank's label
        sign = -sign                                        # credit in bank added to ledger
    return sign                                             

def split_rows(rows):
    col = {name: headers.index(name) for name in headers}   # column name -> position in each row
    unmatched_bank, unmatched_ledger, mismatch = [], [], [] # the three lists we return
    for r in rows:                                          # go through every row marked No
        if r[col["Notes"]] == "amount mismatch":            # both entries exist but the amounts differ
            mismatch.append(r)                              # keep the pair together, the difference is adjusted
        else:                                               # no match, low score, or anything else the human marked No
            if r[col["Bank Amount"]] != "-":                # the row has a bank entry
                unmatched_bank.append(r)                    # break the pair: bank entry becomes unmatched
            if r[col["Ledger Amount"]] != "-":              # the row has a ledger entry
                unmatched_ledger.append(r)                  # break the pair: ledger entry becomes unmatched
    return unmatched_bank, unmatched_ledger, mismatch       

def read_review_sheet(path="review.xlsx"):
    ws = load_workbook(path, data_only=True)["Review"]
    match_col = headers.index("Match?")                    # use the header list, so no magic column number
    rows = list(ws.iter_rows(min_row=2, values_only=True))  # min_row=2 skips the header row
    return [r for r in rows if r[match_col] == "No"]       # only rows still marked No go into the statement

def read_balances(path="review.xlsx"):
    ws = load_workbook(path, data_only=True)["Balances"]    # the second sheet written by write_review_sheet
    return ws["B1"].value, ws["B2"].value                   # B1 = bank balance, B2 = ledger balance

money = '#,##0.00;(#,##0.00)'   # negatives show in brackets, like the template's (X)

def start_statement(ws, bank_balance, ledger_balance, company, as_at):
    ws["A1"], ws["A2"], ws["A3"] = company.upper(), "BANK RECONCILIATION", f"AS AT {as_at.upper()}"
    ws["A5"], ws["B5"], ws["D5"], ws["E5"] = "BANK STATEMENT", "$", "CASH BOOK", "$"
    ws["A6"], ws["B6"] = "UNADJUSTED CLOSING BALANCE:", bank_balance
    ws["D6"], ws["E6"] = "UNADJUSTED CLOSING BALANCE:", ledger_balance
    for cell in (ws["B6"], ws["E6"]):
        cell.number_format = money

# function that writes rows in the sheet
def write_lines(ws, lines, label_col, value_col, start_row=7):
    for offset, (label, value) in enumerate(lines):
        row = start_row + offset                                  # one line per row, under the unadjusted balance
        ws.cell(row=row, column=label_col, value=label)
        ws.cell(row=row, column=value_col, value=value).number_format = money   # negatives show in brackets
    return start_row + len(lines) - 1                             # last row used (6 if there are no lines)

def build_lines(unmatched_bank, unmatched_ledger, mismatch):     # "other" removed
    c = {name: headers.index(name) for name in headers}          # column name -> position in the row
    def line(desc, value):
        return (f"{'ADD' if value >= 0 else 'DEDUCT'} {desc}".upper(), value)   # label follows the sign, like the template

    # ledger entries sit on the bank side
    bank_lines = [line(r[c["Ledger Description"]], row_sign(r) * r[c["Ledger Amount"]]) for r in unmatched_ledger]

    # bank entries sit on the cash book side
    cash_lines = [line(r[c["Bank Description"]], row_sign(r) * r[c["Bank Amount"]]) for r in unmatched_bank]

    # amount difference assumed to be a ledger error and adjusted on the cash book side
    cash_lines += [line("error: " + r[c["Bank Description"]], row_sign(r) * (r[c["Bank Amount"]] - r[c["Ledger Amount"]])) for r in mismatch]
    return bank_lines, cash_lines                                # the two lists main() writes into the sheet

def add_totals(ws, last_bank_row, last_cash_row):
    adj = max(last_bank_row, last_cash_row) + 2          # one shared row, so both adjusted balances line up
    for label_col, val_col in ((1, "B"), (4, "E")):      # bank side (A, B), then cash book side (D, E)
        ws.cell(row=adj, column=label_col, value="ADJUSTED CLOSING BALANCE:").font = Font(bold=True)
        cell = ws[f"{val_col}{adj}"]
        cell.value = f"={val_col}6+SUM({val_col}7:{val_col}{adj - 1})"   # unadjusted balance + all signed lines above
        cell.number_format, cell.font = money, Font(bold=True)
    return adj

def add_check(ws, adj):
    row = adj + 2                                       
    ws[f"A{row}"] = "UNRECONCILED AMOUNT:"               # label, bottom left 
    ws[f"B{row}"] = f"=B{adj}-E{adj}"                    # adjusted bank balance minus adjusted cash book balance
    ws[f"B{row}"].number_format = money                  # same bracket format as the other amounts

    # text next to the amount; ROUND to cents so float noise doesn't show as "Not reconciled"
    ws[f"D{row}"] = f'=IF(ROUND(B{row},2)=0,"Reconciled","Not reconciled")'

    green = PatternFill("solid", bgColor="C6EFCE")       # light green fill for the reconciled case
    red = PatternFill("solid", bgColor="FFC7CE")         # light red fill for the not reconciled case
    # conditional formatting = Excel recolours these cells by itself whenever the formula result changes

    # green when the unreconciled amount rounds to 0
    ws.conditional_formatting.add(f"B{row}:D{row}", FormulaRule(formula=[f"ROUND($B${row},2)=0"], fill=green))
    # red when it doesn't
    ws.conditional_formatting.add(f"B{row}:D{row}", FormulaRule(formula=[f"ROUND($B${row},2)<>0"], fill=red))

def style_statement(ws, adj):
    fill = lambda c: PatternFill("solid", fgColor=c)
    thin = Side(style="thin", color="BFBFBF")
    cols = ("A", "B", "D", "E")                       # C is the spacer column

    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", (52, 16, 3, 52, 16)):
        ws.column_dimensions[col].width = w

    # title block: navy band
    for r, size in ((1, 16), (2, 13), (3, 11)):
        for col in "ABCDE":
            ws[f"{col}{r}"].fill = fill("1F3864")
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"].font = Font(bold=True, size=size, color="FFFFFF")
        ws[f"A{r}"].alignment = Alignment(horizontal="center")

    # side headers
    for col in cols:
        ws[f"{col}5"].fill = fill("2F5597")
        ws[f"{col}5"].font = Font(bold=True, color="FFFFFF")
    for col in ("B", "E"):
        ws[f"{col}5"].alignment = Alignment(horizontal="right")

    # unadjusted balance row
    for col in cols:
        ws[f"{col}6"].fill = fill("D9E2F3")
        ws[f"{col}6"].font = Font(bold=True)

    # adjustment lines: green additions, red brackets for deductions
    for row in range(7, adj - 1):
        for col in cols:
            ws[f"{col}{row}"].border = Border(bottom=thin)
        for col in ("B", "E"):
            ws[f"{col}{row}"].number_format = '[Color10]#,##0.00;[Red](#,##0.00)'

    # adjusted balance row
    for col in cols:
        ws[f"{col}{adj}"].fill = fill("D9E2F3")
        ws[f"{col}{adj}"].border = Border(top=thin, bottom=Side(style="double"))

    # unreconciled row label (the green/red fill comes from your conditional formatting)
    ws[f"A{adj + 2}"].font = Font(bold=True)
    ws[f"B{adj + 2}"].font = Font(bold=True)
    ws[f"D{adj + 2}"].font = Font(bold=True)

def run_pipeline(path):
    rows = read_review_sheet(path)
    unmatched_bank, unmatched_ledger, mismatch = split_rows(rows)
    as_at = date.today().strftime("%d %B %Y")
    bank_balance, ledger_balance = read_balances(path)
    company = "Company"

    wb = load_workbook(path)                          # normal load, so Review/Balances stay intact
    if "Reconciliation" in wb.sheetnames:             # re-run: replace the old statement
        del wb["Reconciliation"]
    ws = wb.create_sheet("Reconciliation")

    start_statement(ws, bank_balance, ledger_balance, company, as_at)
    bank_lines, cash_lines = build_lines(unmatched_bank, unmatched_ledger, mismatch)
    last_bank_row = write_lines(ws, bank_lines, 1, 2)
    last_cash_row = write_lines(ws, cash_lines, 4, 5)
    adj = add_totals(ws, last_bank_row, last_cash_row)
    add_check(ws, adj)
    style_statement(ws, adj)

    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"Close {path} in Excel and run again")

# once approve is clicked (later), all "match?=no" entries should be used to create the bank reconciliation statement 
# automatic trigger when connected with frontend
def main(folder="output"):
    for f in sorted(Path(folder).glob("review_*.xlsx")):   # every review file made by review.py
        if f.name.startswith("~$"):                        # skip Excel's temp lock files
            continue
        print(f"Processing {f.name}")
        run_pipeline(f)
        print("Added Reconciliation sheet to", f.name)


main()